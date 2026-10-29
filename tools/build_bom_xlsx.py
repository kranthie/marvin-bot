"""Build a multi-sheet spreadsheet of the BOM from the CSVs in bom/.

The CSVs stay the source of truth; the spreadsheet is generated (outputs/ is git-ignored). One
sheet per BOM group with line totals as formulas, a summary of totals per group and per phase,
the sourcing table, and the CAD's made-parts and actuator-bolt lists.

Usage:
  pip install -e ".[bom]"
  python tools/build_bom_xlsx.py                 # writes outputs/bom.xlsx
  python tools/build_bom_xlsx.py -o my_bom.xlsx
"""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
BOM_DIR = ROOT / "bom"

# BOM group -> sheet title, in the order the sheets appear
GROUPS = {
    "actuators": "Actuators",
    "hands": "Hands",
    "compute": "Compute & sensors",
    "power": "Power",
    "electronics": "Electronics & wiring",
    "structure": "Structure & materials",
    "fasteners": "Fasteners",
    "tooling": "Tooling",
    "safety": "Safety",
}
# Not in bom.csv: shown on the summary so a replicator sees it, never added to the totals
OPTIONAL = (
    "Optional, not counted: a desktop RTX GPU with 24-32 GB for training and GR00T fine-tuning "
    "(RTX 5090: $1,999 list, about $4,000 or more street price in late 2026). Also not counted: "
    "the desktop CNC mill and 3D printer, shipping, import duties and tariffs."
)

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DCE6F1")
BOLD = Font(bold=True)
LINK = Font(color="0563C1", underline="single")
TOP = Border(top=Side(style="thin"))
USD = '"$"#,##0.00'
USD0 = '"$"#,##0'
WRAP = Alignment(wrap_text=True, vertical="top")

# (header, csv field, width, number format)
LINE_COLUMNS = [
    ("ID", "id", 18, None),
    ("Item", "item", 48, None),
    ("Model", "model", 9, None),
    ("Phase", "phase", 7, None),
    ("Status", "status", 9, None),
    ("Spare", "spare", 7, None),
    ("Qty", "qty", 6, "0"),
    ("Unit USD", "unit_usd", 11, USD),
    ("Line USD", None, 12, USD),
    ("Price basis", "price_basis", 44, None),
    ("Source", "source", 30, None),
    ("Notes", "notes", 60, None),
]
COL = {h: get_column_letter(i + 1) for i, (h, *_) in enumerate(LINE_COLUMNS)}


def read_csv(name: str) -> list[dict]:
    with (BOM_DIR / name).open(newline="") as f:
        return list(csv.DictReader(f))


def number(value: str):
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def header(ws, titles: list[str], widths: list[int]) -> None:
    for i, (title, width) in enumerate(zip(titles, widths, strict=True), start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 20


def link(cell, url: str) -> None:
    if url.startswith("http"):
        cell.hyperlink, cell.font = url, LINK


def group_sheet(wb, title: str, rows: list[dict]) -> tuple[str, int, int]:
    """Write one group's lines; return (sheet, last line row, total row)."""
    ws = wb.create_sheet(title)
    header(ws, [h for h, *_ in LINE_COLUMNS], [w for _, _, w, _ in LINE_COLUMNS])
    rows = sorted(rows, key=lambda r: (r["phase"], r["spare"], r["item"], r["id"]))
    for n, r in enumerate(rows, start=2):
        for i, (_, field, _, fmt) in enumerate(LINE_COLUMNS, start=1):
            if field is None:
                value = f"={COL['Qty']}{n}*{COL['Unit USD']}{n}"
            elif field in ("qty", "unit_usd"):
                value = number(r[field])
            else:
                value = r[field] or None
            c = ws.cell(row=n, column=i, value=value)
            if fmt:
                c.number_format = fmt
            if field in ("item", "price_basis", "notes"):
                c.alignment = WRAP
            elif field == "source" and value:
                link(c, value)
    last = len(rows) + 1
    total = last + 2
    ws[f"B{total}"] = f"Total {title.lower()}"
    ws[f"{COL['Line USD']}{total}"] = f"=SUM({COL['Line USD']}2:{COL['Line USD']}{last})"
    for i in range(1, len(LINE_COLUMNS) + 1):
        c = ws.cell(row=total, column=i)
        c.fill, c.font, c.border = TOTAL_FILL, BOLD, TOP
    ws[f"{COL['Line USD']}{total}"].number_format = USD0
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LINE_COLUMNS))}{last}"
    return title, last, total


def summary_sheet(ws, sheets: list[tuple[str, int, int]], phases: list[str]) -> None:
    ws.title = "Summary"
    ws["A1"] = "Marvin Bot bill of materials"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Generated from bom/*.csv by tools/build_bom_xlsx.py; the CSVs are the source of truth. "
        "Prices are indicative (see each line's price basis and the Sourcing sheet)."
    )
    for col, width in zip("ABCDE", (28, 9, 14, 14, 14), strict=True):
        ws.column_dimensions[col].width = width

    def table_header(row: int, titles: list[str]) -> None:
        for i, t in enumerate(titles, start=1):
            c = ws.cell(row=row, column=i, value=t)
            c.fill, c.font = HEADER_FILL, HEADER_FONT

    def rng(sheet: str, last: int, col: str) -> str:
        return f"'{sheet}'!${COL[col]}$2:${COL[col]}${last}"

    usd, status, phase = "Line USD", "Status", "Phase"
    table_header(4, ["Group", "Lines", "Total USD", "Owned USD", "To buy USD"])
    r = 5
    for sheet, last, total in sheets:
        ws.cell(row=r, column=1, value=sheet)
        ws.cell(row=r, column=2, value=last - 1)
        ws.cell(row=r, column=3, value=f"='{sheet}'!{COL[usd]}{total}")
        own = f'=SUMIF({rng(sheet, last, status)},"owned",{rng(sheet, last, usd)})'
        ws.cell(row=r, column=4, value=own)
        ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
        r += 1
    group_total = r
    ws.cell(row=r, column=1, value="Total (full cost to replicate)")
    ws.cell(row=r, column=2, value=f"=SUM(B5:B{r - 1})")
    for col in "CDE":
        ws[f"{col}{r}"] = f"=SUM({col}5:{col}{r - 1})"
    for c in ws[r][:5]:
        c.fill, c.font, c.border = TOTAL_FILL, BOLD, TOP
    for row in ws.iter_rows(min_row=5, max_row=r, min_col=3, max_col=5):
        for c in row:
            c.number_format = USD0

    r += 2
    table_header(r, ["Phase", "", "Total USD", "Owned USD", "To buy USD"])
    first = r + 1
    for p in phases:
        r += 1
        ws.cell(row=r, column=1, value=p)
        parts, owned = [], []
        for sheet, last, _ in sheets:
            parts.append(f'SUMIF({rng(sheet, last, phase)},"{p}",{rng(sheet, last, usd)})')
            owned.append(
                f'SUMIFS({rng(sheet, last, usd)},{rng(sheet, last, phase)},"{p}",'
                f'{rng(sheet, last, status)},"owned")'
            )
        ws.cell(row=r, column=3, value="=" + "+".join(parts))
        ws.cell(row=r, column=4, value="=" + "+".join(owned))
        ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
    r += 1
    ws.cell(row=r, column=1, value="Total")
    for col in "CDE":
        ws[f"{col}{r}"] = f"=SUM({col}{first}:{col}{r - 1})"
    for c in ws[r][:5]:
        c.fill, c.font, c.border = TOTAL_FILL, BOLD, TOP
    for row in ws.iter_rows(min_row=first, max_row=r, min_col=3, max_col=5):
        for c in row:
            c.number_format = USD0
    ws.cell(row=r + 1, column=1, value=f"Check: equals the group total (C{group_total})")
    ws.cell(row=r + 1, column=3, value=f'=IF(ROUND(C{r}-C{group_total},2)=0,"OK","MISMATCH")')

    r += 3
    ws.cell(row=r, column=1, value=OPTIONAL).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 62
    ws.freeze_panes = "A4"


def table_sheet(wb, title: str, rows: list[dict], widths: dict[str, int], sort_key) -> None:
    ws = wb.create_sheet(title)
    fields = list(rows[0])
    header(ws, fields, [widths.get(f, 12) for f in fields])
    for n, r in enumerate(sorted(rows, key=sort_key), start=2):
        for i, f in enumerate(fields, start=1):
            value = number(r[f]) if f in ("qty", "unit_usd", "mass_g", "stock_mm") else r[f]
            c = ws.cell(row=n, column=i, value=value if value != "" else None)
            if f == "unit_usd" and value != "":
                c.number_format = USD
            if f == "url" and value:
                link(c, value)
            if f in ("notes", "part", "tiers", "availability"):
                c.alignment = WRAP
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(rows) + 1}"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("-o", "--out", type=Path, default=ROOT / "outputs" / "bom.xlsx")
    args = ap.parse_args()

    bom = read_csv("bom.csv")
    unknown = {r["group"] for r in bom} - set(GROUPS)
    if unknown:
        raise SystemExit(f"groups without a sheet: {sorted(unknown)}")

    wb = Workbook()
    summary = wb.active
    sheets = [
        group_sheet(wb, title, [r for r in bom if r["group"] == g]) for g, title in GROUPS.items()
    ]
    summary_sheet(summary, sheets, sorted({r["phase"] for r in bom}))
    table_sheet(
        wb,
        "Sourcing",
        read_csv("sourcing.csv"),
        {"id": 18, "vendor": 30, "kind": 16, "url": 50, "tiers": 22, "availability": 26},
        lambda r: (r["id"], r["vendor"]),
    )
    table_sheet(
        wb,
        "Made parts (CAD)",
        read_csv("build_parts.csv"),
        {"part": 34, "form": 13, "size_mm": 18, "notes": 60},
        lambda r: (r["process"], r["material"], r["part"]),
    )
    table_sheet(
        wb,
        "Actuator bolts (CAD)",
        read_csv("build_fasteners.csv"),
        {"notes": 60},
        lambda r: (r["actuator"], r["face"]),
    )
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out} ({len(bom)} BOM lines on {len(sheets)} group sheets)")


if __name__ == "__main__":
    main()
