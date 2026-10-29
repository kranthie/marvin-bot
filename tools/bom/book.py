"""Maintain the BOM workbook, bom/bom.xlsx: the bill of materials' single source of truth.

The workbook has a Summary, one sheet per group of bought items, the Sourcing table, and two
sheets of made parts and actuator bolts that `python -m cad.export` keeps up to date from the
CAD. This tool rebuilds the layout, sorting, formatting, formulas and dropdowns of those sheets
from the rows already in the workbook, so hand edits and scripted edits come out the same. Sheets
it does not manage are left alone, and so are the columns a builder fills in (status, build notes).

Usage:
  python tools/bom/book.py            # normalise bom/bom.xlsx in place (after editing it by hand)
  python tools/bom/book.py --check    # exit 1 if it has errors or is not normalised (CI)
  python tools/bom/book.py set cells unit_usd=9.95 "price_basis=dealer 2026-11-02"
  python tools/bom/book.py add usb-hub group=electronics "item=USB 3 hub" qty=1 unit_usd=30
  python tools/bom/book.py source cells "vendor=Some dealer" url=https://... unit_usd=9.95
  python tools/bom/book.py remove usb-hub
Each edit normalises the workbook afterwards. `source` dates the row today unless checked= is
given. Fields are the column names below (id, group, item, model, status, qty, spares, ...).
Each item has one line: qty is what the robot uses, spares the extra units bought.
"""

import argparse
import datetime
import io
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "bom" / "bom.xlsx"

# BOM group -> sheet title, in the order the sheets appear
GROUPS = {
    "actuators": "Actuators",
    "hands": "Hands",
    "compute": "Compute & sensors",
    "power": "Power",
    "electronics": "Electronics & wiring",
    "structure": "Structure & materials",
    "fasteners": "Fasteners",
    "tooling": "Tooling",
    "safety": "Safety",
}
SUMMARY, SOURCING, PARTS, BOLTS = (
    "Summary",
    "Sourcing",
    "Made parts (CAD)",
    "Actuator bolts (CAD)",
)
STATUSES = ("to buy", "ordered", "owned")
PART_STATUSES = ("to make", "ordered", "made")
# Not in the BOM: shown on the summary so a replicator sees it, never added to the totals
OPTIONAL = (
    "Optional, not counted: a desktop RTX GPU with 24-32 GB for training and GR00T fine-tuning "
    "(RTX 5090: $1,999 list, about $4,000 or more street price in late 2026). Also not counted: "
    "the desktop CNC mill and 3D printer, shipping, import duties and tariffs."
)

USD = '"$"#,##0.00'
USD0 = '"$"#,##0'
# (field, header, width, number format); line_usd is a formula, not data
LINE_COLUMNS = [
    ("id", "ID", 18, None),
    ("item", "Item", 48, None),
    ("model", "Model", 9, None),
    ("status", "Status", 9, None),
    ("qty", "Qty", 6, "0"),
    ("spares", "Spares", 7, "0"),
    ("unit_usd", "Unit USD", 11, USD),
    ("line_usd", "Line USD", 12, USD),
    ("price_basis", "Price basis", 44, None),
    ("source", "Source", 30, None),
    ("notes", "Notes", 60, None),
]
SOURCING_COLUMNS = [
    ("id", "ID", 18, None),
    ("vendor", "Vendor", 30, None),
    ("kind", "Kind", 16, None),
    ("url", "URL", 50, None),
    ("unit_usd", "Unit USD", 11, USD),
    ("tiers", "Tiers", 22, None),
    ("availability", "Availability", 26, None),
    ("checked", "Checked", 11, None),
    ("confirmed", "Confirmed", 10, None),
    ("notes", "Notes", 60, None),
]
# made parts and bolts: the CAD writes the first columns; status and build notes are the builder's
PARTS_COLUMNS = [
    ("part", "Part", 34, None),
    ("qty", "Qty", 6, "0"),
    ("process", "Process", 9, None),
    ("material", "Material", 9, None),
    ("form", "Form", 13, None),
    ("stock_mm", "Stock mm", 9, None),
    ("mass_g", "Mass g", 9, "0.0"),
    ("size_mm", "Size mm", 18, None),
    ("notes", "Notes", 60, None),
    ("status", "Status", 10, None),
    ("build_notes", "Build notes", 40, None),
]
PARTS_GENERATED = [f for f, *_ in PARTS_COLUMNS[:9]]
BOLTS_COLUMNS = [
    ("actuator", "Actuator", 10, None),
    ("face", "Face", 8, None),
    ("thread", "Thread", 8, None),
    ("qty", "Qty", 6, "0"),
    ("notes", "Notes", 60, None),
    ("status", "Status", 9, None),
]
BOLTS_GENERATED = [f for f, *_ in BOLTS_COLUMNS[:5]]
WRAP_FIELDS = {"item", "price_basis", "notes", "tiers", "availability", "build_notes"}
LINK_FIELDS = {"source", "url"}

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DCE6F1")
BOLD = Font(bold=True)
LINK = Font(color="0563C1", underline="single")
TOP = Border(top=Side(style="thin"))
WRAP = Alignment(wrap_text=True, vertical="top")


# ---- reading -------------------------------------------------------------------------------


def read_table(ws, columns) -> list[dict]:
    """Rows under the header, down to the first row with an empty first column."""
    names = {h: f for f, h, *_ in columns} | {f: f for f, *_ in columns}
    head = [names.get(c.value) for c in ws[1]]
    rows = []
    for cells in ws.iter_rows(min_row=2, values_only=True):
        if cells[0] in (None, ""):
            break
        rows.append({f: v for f, v in zip(head, cells, strict=False) if f and f != "line_usd"})
    return rows


def read_book(path: Path = BOOK) -> dict:
    wb = load_workbook(path)
    return read_wb(wb)


def read_wb(wb) -> dict:
    lines = []
    for group, title in GROUPS.items():
        if title in wb.sheetnames:
            lines += [dict(r, group=group) for r in read_table(wb[title], LINE_COLUMNS)]
    get = {
        SOURCING: SOURCING_COLUMNS,
        PARTS: PARTS_COLUMNS,
        BOLTS: BOLTS_COLUMNS,
    }
    sourcing, parts, bolts = (
        read_table(wb[t], c) if t in wb.sheetnames else [] for t, c in get.items()
    )
    return dict(lines=lines, sourcing=sourcing, parts=parts, bolts=bolts)


def validate(data: dict) -> list[str]:
    errors, ids = [], set()
    for r in data["lines"]:
        i = r.get("id")
        if i in ids:
            errors.append(f"duplicate id {i}")
        ids.add(i)
        if r.get("group") not in GROUPS:
            errors.append(f"{i}: unknown group {r.get('group')}")
        if r.get("status") not in (None, *STATUSES):
            errors.append(f"{i}: status must be empty or one of {', '.join(STATUSES)}")
        if not all(isinstance(r.get(f), int) and r[f] >= 0 for f in ("qty", "spares")):
            errors.append(f"{i}: qty and spares must be whole numbers (0 or more)")
        if not isinstance(r.get("unit_usd"), int | float):
            errors.append(f"{i}: unit_usd must be a number")
        if not r.get("item"):
            errors.append(f"{i}: no item")
    for r in data["sourcing"]:
        if r.get("id") not in ids:
            errors.append(f"Sourcing: {r.get('id')} is not a BOM line")
        if r.get("confirmed") not in ("yes", "no"):
            errors.append(f"Sourcing: {r.get('id')}: confirmed must be yes or no")
    for r in data["parts"]:
        if r.get("status") not in (None, *PART_STATUSES):
            errors.append(f"{PARTS}: {r.get('part')}: unknown status {r.get('status')}")
    for r in data["bolts"]:
        if r.get("status") not in (None, *STATUSES):
            errors.append(f"{BOLTS}: {r.get('actuator')}: unknown status {r.get('status')}")
    return errors


# ---- writing -------------------------------------------------------------------------------


def blank(v):
    return None if v == "" else v


def header(ws, columns) -> None:
    for i, (_, title, width, _) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 20


def write_rows(ws, columns, rows) -> int:
    """Header and rows; return the last row's number."""
    header(ws, columns)
    for n, r in enumerate(rows, start=2):
        for i, (field, _, _, fmt) in enumerate(columns, start=1):
            c = ws.cell(row=n, column=i, value=blank(r.get(field)))
            if fmt:
                c.number_format = fmt
            if field in WRAP_FIELDS:
                c.alignment = WRAP
            elif field in LINK_FIELDS and isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink, c.font = c.value, LINK
    last = len(rows) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"
    return last


def dropdown(ws, columns, field, options, last) -> None:
    col = get_column_letter([f for f, *_ in columns].index(field) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = f"Choose one of: {', '.join(options)}"
    dv.add(f"{col}2:{col}{max(last, 2)}")
    ws.add_data_validation(dv)


COL = {f: get_column_letter(i + 1) for i, (f, *_) in enumerate(LINE_COLUMNS)}


def line_value(r: dict) -> float:
    """Cost of a line: the units used plus the spares."""
    try:
        return (r["qty"] + r["spares"]) * r["unit_usd"]
    except TypeError:
        return 0


def line_sheet(ws, title: str, rows: list[dict], cache: dict) -> tuple[int, int]:
    """One group's lines, a formula per line and a total; return (last line row, total row)."""
    rows = sorted(rows, key=lambda r: (r["item"], r["id"]))
    last = write_rows(ws, LINE_COLUMNS, rows)
    for n, r in enumerate(rows, start=2):
        units = f"({COL['qty']}{n}+{COL['spares']}{n})"
        ws[f"{COL['line_usd']}{n}"] = f"={units}*{COL['unit_usd']}{n}"
        cache[f"{COL['line_usd']}{n}"] = line_value(r)
    total = last + 2
    ws[f"B{total}"] = f"Total {title.lower()}"
    ws[f"{COL['line_usd']}{total}"] = f"=SUM({COL['line_usd']}2:{COL['line_usd']}{last})"
    cache[f"{COL['line_usd']}{total}"] = sum(line_value(r) for r in rows)
    for i in range(1, len(LINE_COLUMNS) + 1):
        c = ws.cell(row=total, column=i)
        c.fill, c.font, c.border = TOTAL_FILL, BOLD, TOP
    ws[f"{COL['line_usd']}{total}"].number_format = USD0
    ws.freeze_panes = "C2"
    dropdown(ws, LINE_COLUMNS, "status", STATUSES, last)
    return last, total


def summary_sheet(ws, sheets: list[tuple[str, int, int]], groups: list[list[dict]], cache) -> None:
    """Totals per group; cache gets each formula's value, for viewers that do not compute
    formulas."""
    ws["A1"] = "Marvin Bot bill of materials"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "This workbook is the BOM's source of truth; copy it to track your own build. Set a "
        "line's status to owned or ordered as you go (empty counts as to buy) and the totals "
        "follow. After editing, run "
        "python tools/bom/book.py to tidy it. Prices are indicative: see each line's price "
        "basis and the Sourcing sheet."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 48
    for col, width in zip("ABCDEF", (30, 9, 13, 13, 13, 13), strict=True):
        ws.column_dimensions[col].width = width
    titles = ["Total USD", "Owned USD", "Ordered USD", "To buy USD"]

    def rng(sheet: str, last: int, field: str) -> str:
        return f"'{sheet}'!${COL[field]}$2:${COL[field]}${last}"

    def head(row: int, first: str, second: str) -> None:
        for i, t in enumerate([first, second, *titles], start=1):
            c = ws.cell(row=row, column=i, value=t or None)
            c.fill, c.font = HEADER_FILL, HEADER_FONT

    def sums(row: int, crit: list[tuple[str, str]], sheets_, rows_) -> None:
        for col, extra in ((3, []), (4, ["owned"]), (5, ["ordered"])):
            c = crit + [("status", s) for s in extra]
            ws.cell(row=row, column=col, value=sum_expr(c, sheets_))
            match = [r for r in rows_ if all(r.get(f) == v for f, v in c)]
            cache[f"{get_column_letter(col)}{row}"] = sum(line_value(r) for r in match)
        # to buy: everything not owned or ordered, so an empty status counts as to buy
        ws.cell(row=row, column=6, value=f"=C{row}-D{row}-E{row}")
        cache[f"F{row}"] = cache[f"C{row}"] - cache[f"D{row}"] - cache[f"E{row}"]

    def sum_expr(crit, sheets_) -> str:
        if not crit:
            return f"='{sheets_[0][0]}'!{COL['line_usd']}{sheets_[0][2]}"
        parts = []
        for sheet, last, _ in sheets_:
            args = ",".join(f'{rng(sheet, last, f)},"{v}"' for f, v in crit)
            parts.append(f"SUMIFS({rng(sheet, last, 'line_usd')},{args})")
        return "=" + "+".join(parts)

    def total_row(row: int, label: str, first: int) -> None:
        ws.cell(row=row, column=1, value=label)
        for col in "CDEF":
            ws[f"{col}{row}"] = f"=SUM({col}{first}:{col}{row - 1})"
            cache[f"{col}{row}"] = sum(cache[f"{col}{n}"] for n in range(first, row))
        for c in ws[row][:6]:
            c.fill, c.font, c.border = TOTAL_FILL, BOLD, TOP

    head(4, "Group", "Lines")
    r = 5
    for s, rows in zip(sheets, groups, strict=True):
        ws.cell(row=r, column=1, value=s[0])
        ws.cell(row=r, column=2, value=f"=COUNTA('{s[0]}'!A2:A{s[1]})")
        cache[f"B{r}"] = len(rows)
        sums(r, [], [s], rows)
        r += 1
    total_row(r, "Total (full cost to replicate)", 5)
    ws[f"B{r}"] = f"=SUM(B5:B{r - 1})"
    cache[f"B{r}"] = sum(len(rows) for rows in groups)
    ws[f"B{r}"].fill, ws[f"B{r}"].font, ws[f"B{r}"].border = TOTAL_FILL, BOLD, TOP
    for row in ws.iter_rows(min_row=5, max_row=r, min_col=3, max_col=6):
        for c in row:
            c.number_format = USD0
    r += 2
    ws.cell(row=r, column=1, value=OPTIONAL).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 48
    ws.freeze_panes = "A4"


def normalise(wb, data: dict) -> dict:
    """Rebuild the managed sheets of wb from data; leave other sheets where they are. Return the
    formulas' values per sheet and cell, for to_bytes to store with them."""
    managed = [SUMMARY, *GROUPS.values(), SOURCING, PARTS, BOLTS]
    for title in managed:
        if title in wb.sheetnames:
            wb.remove(wb[title])
    ws = {t: wb.create_sheet(t, i) for i, t in enumerate(managed)}
    sheets, groups, cache = [], [], {t: {} for t in managed}
    for group, title in GROUPS.items():
        rows = [r for r in data["lines"] if r["group"] == group]
        last, total = line_sheet(ws[title], title, rows, cache[title])
        sheets.append((title, last, total))
        groups.append(rows)
    summary_sheet(ws[SUMMARY], sheets, groups, cache[SUMMARY])
    srcs = sorted(data["sourcing"], key=lambda r: (r["id"], r.get("vendor") or ""))
    write_rows(ws[SOURCING], SOURCING_COLUMNS, srcs)
    dropdown(ws[SOURCING], SOURCING_COLUMNS, "confirmed", ("yes", "no"), len(srcs) + 1)
    parts = sorted(data["parts"], key=lambda r: (r["process"], r["material"], r["part"]))
    write_rows(ws[PARTS], PARTS_COLUMNS, parts)
    dropdown(ws[PARTS], PARTS_COLUMNS, "status", PART_STATUSES, len(parts) + 1)
    bolts = sorted(data["bolts"], key=lambda r: (r["actuator"], r["face"]))
    write_rows(ws[BOLTS], BOLTS_COLUMNS, bolts)
    dropdown(ws[BOLTS], BOLTS_COLUMNS, "status", STATUSES, len(bolts) + 1)
    wb.active = 0
    for w in wb.worksheets:
        w.sheet_view.tabSelected = w.title == SUMMARY
    wb.calculation.fullCalcOnLoad = True
    return cache


def fingerprint(wb) -> list:
    """What --check compares: values, formulas and the formatting this tool sets."""
    out = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and not c.has_style:
                    continue
                link = c.hyperlink.target if c.hyperlink else None
                color = c.font.color.rgb if c.font and c.font.color else None
                cells.append(
                    (c.coordinate, c.value, c.number_format, c.font.b, color)
                    + (c.fill.fgColor.rgb, link, c.alignment.wrap_text)
                )
        widths = {k: d.width for k, d in ws.column_dimensions.items() if d.width}
        dvs = sorted((str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation)
        merged = sorted(str(m) for m in ws.merged_cells.ranges)
        out.append((ws.title, cells, widths, ws.freeze_panes, dvs, merged, ws.auto_filter.ref))
    return out


def computed(raw: bytes) -> list:
    """Every cell's stored value, formulas' cached results included (what a preview shows)."""
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    return [[c.value for row in ws.iter_rows() for c in row] for ws in wb.worksheets]


def to_bytes(wb, cache: dict | None = None) -> bytes:
    """Save wb; store each formula's value from cache next to it, as a spreadsheet program does,
    so viewers that do not compute formulas (previews, web views) still show the totals."""
    buf = io.BytesIO()
    wb.save(buf)
    if not cache:
        return buf.getvalue()
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    rels = src.read("xl/_rels/workbook.xml.rels").decode()
    target = {}
    for rel in re.findall(r"<Relationship [^>]*>", rels):
        rid, tgt = re.search(r'Id="([^"]+)"', rel), re.search(r'Target="/?(?:xl/)?([^"]+)"', rel)
        if rid and tgt:
            target[rid.group(1)] = tgt.group(1)
    books = src.read("xl/workbook.xml").decode()
    files = {
        "xl/" + target[rid]: cache.get(name.replace("&amp;", "&"), {})
        for name, rid in re.findall(r'<sheet name="([^"]+)"[^>]*r:id="([^"]+)"', books)
    }

    def fill(values: dict):
        def put(m):
            v = values.get(m.group(1))
            if v is None:
                return m.group(0)
            if isinstance(v, str):
                return f'<c r="{m.group(1)}"{m.group(2)} t="str"><f>{m.group(3)}</f><v>{v}</v>'
            v = round(v, 2) if isinstance(v, float) else v
            return f'<c r="{m.group(1)}"{m.group(2)}><f>{m.group(3)}</f><v>{v!r}</v>'

        return put

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if files.get(info.filename):
                xml = data.decode()
                pattern = r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/>'
                data = re.sub(pattern, fill(files[info.filename]), xml).encode()
            dst.writestr(info, data)
    return out.getvalue()


def normalised_bytes(path: Path, data: dict | None = None) -> tuple[bytes, list, list]:
    """(normalised workbook, its fingerprint, the file's fingerprint)."""
    wb = load_workbook(path)
    current = [fingerprint(wb), computed(path.read_bytes())]
    raw = to_bytes(wb, normalise(wb, read_wb(wb) if data is None else data))
    return raw, [fingerprint(load_workbook(io.BytesIO(raw))), computed(raw)], current


def save(data: dict | None = None, path: Path = BOOK) -> bool:
    """Normalise the workbook (with data instead of its own rows, if given); write if it changed."""
    raw, new, current = normalised_bytes(path, data)
    if new == current:
        return False
    path.write_bytes(raw)
    return True


def check(path: Path = BOOK) -> list[str]:
    data = read_book(path)
    errors = validate(data)
    _, new, current = normalised_bytes(path)
    if new != current:
        pairs = zip(new[0], current[0], strict=False)
        sheets = [a[0] for a, b in pairs if a != b] or ["stored formula values or sheet order"]
        errors.append(
            f"{path.name} is not normalised ({', '.join(sheets)}): run python tools/bom/book.py"
        )
    return errors


# ---- the CAD's sheets ----------------------------------------------------------------------


def generated_stale(data: dict, parts: list[dict], bolts: list[dict]) -> list[str]:
    """Names of the CAD sheets whose generated columns differ from the CAD export's rows."""

    def norm(v):  # 25 and 25.0 are the same mass
        return float(v) if isinstance(v, int | float) else blank(v)

    def key(rows, fields):
        return sorted(tuple(str(norm(r.get(f))) for f in fields) for r in rows)

    stale = []
    if key(data["parts"], PARTS_GENERATED) != key(parts, PARTS_GENERATED):
        stale.append(PARTS)
    if key(data["bolts"], BOLTS_GENERATED) != key(bolts, BOLTS_GENERATED):
        stale.append(BOLTS)
    return stale


def merge_generated(data: dict, parts: list[dict], bolts: list[dict]) -> dict:
    """The CAD's new rows with the builder's columns carried over from matching old rows."""
    old_parts = {r["part"]: r for r in data["parts"]}
    old_bolts = {(r["actuator"], r["face"]): r for r in data["bolts"]}
    new_parts = []
    for r in parts:
        old = old_parts.get(r["part"], {})
        extra = {"status": old.get("status"), "build_notes": old.get("build_notes")}
        new_parts.append({k: blank(v) for k, v in r.items()} | extra)
    new_bolts = []
    for r in bolts:
        old = old_bolts.get((r["actuator"], r["face"]), {})
        new_bolts.append(dict(r, status=old.get("status")))
    return dict(data, parts=new_parts, bolts=new_bolts)


# ---- command line --------------------------------------------------------------------------


def parse_fields(pairs: list[str], columns) -> dict:
    known = {f for f, *_ in columns} | {"group"}
    out = {}
    for p in pairs:
        field, sep, value = p.partition("=")
        if not sep or field not in known or field == "line_usd":
            raise SystemExit(f"expected field=value with a field in {sorted(known)}: {p}")
        if field in ("qty", "spares"):
            out[field] = int(value)
        elif field == "unit_usd":
            out[field] = float(value) if "." in value else int(value) if value else None
        else:
            out[field] = value or None
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="exit 1 on errors or if not normalised")
    ap.add_argument("--book", type=Path, default=BOOK)
    ap.add_argument("command", nargs="?", choices=["set", "add", "source", "remove"])
    ap.add_argument("id", nargs="?")
    ap.add_argument("fields", nargs="*", help="field=value")
    args = ap.parse_args()

    if args.check:
        errors = check(args.book)
        for e in errors:
            print(f"ERROR: {e}")
        if not errors:
            print(f"{args.book.name} is valid and normalised")
        return 1 if errors else 0

    data = read_book(args.book)
    lines = {r["id"]: r for r in data["lines"]}
    if args.command and not args.id:
        ap.error("give the line's id")
    if args.command == "set":
        if args.id not in lines:
            raise SystemExit(f"no line {args.id}")
        lines[args.id].update(parse_fields(args.fields, LINE_COLUMNS))
    elif args.command == "add":
        if args.id in lines:
            raise SystemExit(f"line {args.id} exists: use set")
        row = {"spares": 0} | parse_fields(args.fields, LINE_COLUMNS)
        data["lines"].append(dict(row, id=args.id))
    elif args.command == "source":
        if args.id not in lines:
            raise SystemExit(f"no line {args.id}")
        row = {"checked": datetime.date.today().isoformat(), "confirmed": "yes"}
        data["sourcing"].append(row | parse_fields(args.fields, SOURCING_COLUMNS) | {"id": args.id})
    elif args.command == "remove":
        if args.id not in lines:
            raise SystemExit(f"no line {args.id}")
        data["lines"] = [r for r in data["lines"] if r["id"] != args.id]
        data["sourcing"] = [r for r in data["sourcing"] if r["id"] != args.id]
    errors = validate(data)
    for e in errors:
        print(f"ERROR: {e}")
    if errors:
        print("not saved")
        return 1
    changed = save(data, args.book)
    print(f"{args.book.name} {'updated' if changed else 'unchanged'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
